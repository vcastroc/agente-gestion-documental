import re
import sys
import tempfile
import unittest
import shutil
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import modules.database as database
import modules.agent as agent
import modules.email_sender as email_sender
import modules.ollama_client as ollama_client
import agent_worker
from app import _valid_upload, app
from modules.analyzer import calcular_compatibilidad, extraer_habilidades
from modules.file_utils import calcular_hash_archivo
from modules.matching import vacante_compatible_con_usuario
from modules.recruitbot import SALUDO_RESPUESTA, detectar_intencion, responder_recruitbot
from docx import Document
from werkzeug.datastructures import FileStorage


class RecruitAITestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.original_db_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "test.db"
        self.original_uploads_dir = agent.UPLOADS_DIR
        self.original_worker_uploads_dir = agent_worker.UPLOADS_DIR
        agent.UPLOADS_DIR = Path(self.temp_dir.name) / "uploads"
        agent.UPLOADS_DIR.mkdir()
        agent_worker.UPLOADS_DIR = agent.UPLOADS_DIR
        database.inicializar_bd()
        app.config.update(TESTING=True)

    def tearDown(self):
        agent.UPLOADS_DIR = self.original_uploads_dir
        agent_worker.UPLOADS_DIR = self.original_worker_uploads_dir
        database.DB_PATH = self.original_db_path
        self.temp_dir.cleanup()

    def csrf(self, client, path="/login"):
        html = client.get(path).get_data(as_text=True)
        return re.search(r'name="_csrf_token" value="([^"]+)"', html).group(1)

    def login(self, client, email, password):
        return client.post(
            "/login",
            data={"correo": email, "password": password, "_csrf_token": self.csrf(client)},
        )

    def test_skills_are_not_detected_as_substrings(self):
        self.assertEqual(
            extraer_habilidades("Experiencia con JavaScript y PostgreSQL"),
            ["JavaScript", "PostgreSQL"],
        )

    def test_admin_login_and_candidate_role_guard(self):
        admin = app.test_client()
        response = self.login(admin, "admin@demo.com", "admin123")
        self.assertEqual(response.headers["Location"], "/admin/dashboard")

        candidate_id = database.crear_usuario("Candidate", "candidate@example.com", "clave123")
        candidate = app.test_client()
        self.login(candidate, "candidate@example.com", "clave123")
        blocked = candidate.get("/admin/ranking")
        self.assertEqual(blocked.status_code, 302)
        self.assertEqual(blocked.headers["Location"], "/candidato/dashboard")
        self.assertTrue(candidate_id)

    def test_post_requires_csrf(self):
        admin = app.test_client()
        self.login(admin, "admin@demo.com", "admin123")
        response = admin.post("/admin/vacantes", data={"titulo": "Sin token", "empresa": "Demo"})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(any(item["titulo"] == "Sin token" for item in database.obtener_vacantes()))

    def test_recruitbot_widget_is_rendered_and_endpoint_answers(self):
        admin = app.test_client()
        self.login(admin, "admin@demo.com", "admin123")
        page = admin.get("/admin/dashboard")
        self.assertIn("RecruitBot", page.get_data(as_text=True))

        token = self.csrf(admin, "/admin/dashboard")
        response = admin.post(
            "/chatbot",
            json={"message": "Mejor candidato", "_csrf_token": token},
            headers={"X-CSRFToken": token},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.get_json()["answer"])

    def test_recruitbot_requires_csrf_for_json(self):
        client = app.test_client()
        client.get("/login")
        response = client.post("/chatbot", json={"message": "Hola"})
        self.assertEqual(response.status_code, 400)

    def test_recruitbot_greeting_does_not_reuse_previous_ranking_intent(self):
        self.assertEqual(detectar_intencion("Hola"), "saludo")
        with patch("modules.recruitbot.construir_contexto") as context, patch("modules.recruitbot._post_json") as ollama:
            answer = responder_recruitbot(
                "Hola",
                history=[
                    {"role": "user", "content": "Top 5 candidatos"},
                    {"role": "assistant", "content": "Ranking anterior"},
                ],
                role="admin",
                user_id=1,
            )
        self.assertEqual(answer, SALUDO_RESPUESTA)
        context.assert_not_called()
        ollama.assert_not_called()

    def test_recruitbot_rrhh_query_uses_current_message_intent(self):
        self.assertEqual(detectar_intencion("Top 5 candidatos"), "consulta_rrhh")
        with patch.dict("os.environ", {"OLLAMA_ENABLED": "1"}, clear=False):
            with patch("modules.recruitbot._post_json", return_value={"message": {"content": "Respuesta ranking"}}) as ollama:
                answer = responder_recruitbot("Top 5 candidatos", history=[{"role": "user", "content": "Hola"}], role="admin")
        self.assertEqual(answer, "Respuesta ranking")
        prompt = ollama.call_args.args[1]["messages"][0]["content"]
        self.assertIn("INTENCION DEL MENSAJE ACTUAL: consulta_rrhh", prompt)
        self.assertIn("No continues consultas anteriores", prompt)

    def test_invalid_vacancy_cannot_be_applied_to(self):
        database.crear_usuario("Candidate", "candidate@example.com", "clave123")
        candidate = app.test_client()
        self.login(candidate, "candidate@example.com", "clave123")
        token = self.csrf(candidate, "/candidato/vacantes")
        candidate.post("/candidato/postular/999999", data={"_csrf_token": token})
        user = database.obtener_usuario("candidate@example.com")
        self.assertEqual(database.obtener_postulaciones_usuario(user["id"]), [])

    def test_invitation_query_does_not_duplicate_rows(self):
        vacancy = database.obtener_vacantes()[0]
        candidate_id = database.guardar_candidato({"nombre": "Demo", "correo": "demo@example.com"})
        result = {
            "compatibilidad": 100, "estado": "Preseleccionado",
            "habilidades_encontradas": ["Python"], "habilidades_faltantes": [],
        }
        database.guardar_analisis(candidate_id, vacancy["id"], result)
        database.guardar_analisis(candidate_id, vacancy["id"], result)
        database.guardar_correo(candidate_id, "demo@example.com", "Entrevista", "Mensaje", "Enviado", vacancy["id"])
        self.assertEqual(len(database.obtener_invitaciones_usuario("demo@example.com")), 1)

    def test_email_approval_is_scoped_to_active_vacancy(self):
        first = database.obtener_vacantes()[0]
        second_id = database.guardar_vacante({"titulo": "Segunda", "empresa": "Demo"})
        candidate_id = database.guardar_candidato({"nombre": "Demo", "correo": "demo@example.com"})
        first_email = database.guardar_correo(candidate_id, "demo@example.com", "Primera", "Mensaje", "Borrador", first["id"])
        second_email = database.guardar_correo(candidate_id, "demo@example.com", "Segunda", "Mensaje", "Borrador", second_id)
        with patch("modules.email_sender.enviar_correo_smtp") as smtp:
            result = email_sender.enviar_borradores_correo(first["id"])
        first_status = database._rows("SELECT estado FROM correos WHERE id=?", (first_email,))[0]["estado"]
        second_status = database._rows("SELECT estado FROM correos WHERE id=?", (second_email,))[0]["estado"]
        self.assertEqual(result["enviados"], 1)
        smtp.assert_called_once_with("demo@example.com", "Primera", "Mensaje")
        self.assertEqual(first_status, "Enviado")
        self.assertEqual(second_status, "Borrador")

    def test_upload_rejects_spoofed_pdf_extension(self):
        fake_pdf = FileStorage(stream=BytesIO(b"not a real pdf"), filename="cv.pdf")
        self.assertFalse(_valid_upload(fake_pdf))

    def test_weighted_score_considers_experience_and_studies(self):
        vacancy = {
            "habilidades": "Python, SQL", "experiencia": 2, "estudios": "Universitario",
            "peso_habilidades": 70, "peso_experiencia": 20, "peso_estudios": 10,
        }
        text = "Ingeniería de sistemas. Experiencia de 1 año con Python y SQL."
        self.assertEqual(calcular_compatibilidad(text, vacancy), 90.0)

    def test_closed_vacancy_is_not_available_to_candidates(self):
        vacancy_id = database.guardar_vacante({"titulo": "Cerrada", "empresa": "Demo"})
        self.assertTrue(database.actualizar_estado_vacante(vacancy_id, "Cerrada"))
        self.assertNotIn(vacancy_id, [item["id"] for item in database.obtener_vacantes(solo_activas=True)])

    def test_vacancy_area_matches_candidate_career(self):
        matched, _ = vacante_compatible_con_usuario(
            {"carrera": "Ingenieria de Sistemas"},
            {"area": "Tecnologia", "carreras_objetivo": "Software, Informatica"},
        )
        other, _ = vacante_compatible_con_usuario(
            {"carrera": "Contabilidad"},
            {"area": "Tecnologia", "carreras_objetivo": "Software, Informatica"},
        )
        self.assertTrue(matched)
        self.assertFalse(other)

    def test_candidate_vacancies_are_grouped_by_area(self):
        database.guardar_vacante({
            "titulo": "Backend", "empresa": "Demo", "area": "Tecnologia",
            "carreras_objetivo": "Sistemas, Informatica", "habilidades": "Python",
        })
        database.guardar_vacante({
            "titulo": "Asistente Contable", "empresa": "Demo", "area": "Administracion",
            "carreras_objetivo": "Contabilidad", "habilidades": "Excel",
        })
        database.crear_usuario(
            "Sistemas User", "sistemas@example.com", "clave123", "candidato",
            carrera="Ingenieria de Sistemas",
        )
        candidate = app.test_client()
        self.login(candidate, "sistemas@example.com", "clave123")
        response = candidate.get("/candidato/vacantes")
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertLess(html.index("Backend"), html.index("Asistente Contable"))
        self.assertIn("Afin a tu perfil", html)
        self.assertIn("Otras vacantes disponibles", html)

    def test_custom_threshold_controls_agent_state(self):
        from modules.ranking import asignar_estado

        self.assertEqual(asignar_estado(75, 70, 50), "Preseleccionado")
        self.assertEqual(asignar_estado(55, 70, 50), "En observación")

    def test_cv_upload_preanalysis_recommends_active_vacancies_without_applying(self):
        user_id = database.crear_usuario("Candidate", "candidate@example.com", "clave123")
        cv_name = "candidate.docx"
        document = Document()
        document.add_paragraph("Candidate Example")
        document.add_paragraph("candidate@example.com")
        document.add_paragraph("Ingeniería de sistemas con 2 años de experiencia en Python SQL Git Excel Power BI y trabajo en equipo.")
        document.save(agent.UPLOADS_DIR / cv_name)
        database.guardar_cv_usuario(user_id, cv_name)
        result = agent.ejecutar_preanalisis_perfil(user_id)
        self.assertEqual(result["estado"], "Completado")
        self.assertTrue(database.obtener_recomendaciones_usuario(user_id))
        self.assertEqual(database.obtener_postulaciones_usuario(user_id), [])

    def test_worker_processes_cv_preanalysis_from_persistent_queue(self):
        user_id = database.crear_usuario("Worker Candidate", "worker@example.com", "clave123")
        cv_name = "worker_candidate.docx"
        document = Document()
        document.add_paragraph("Worker Candidate")
        document.add_paragraph("worker@example.com")
        document.add_paragraph("Ingeniería de sistemas con 2 años de experiencia en Python SQL Git Excel.")
        document.save(agent.UPLOADS_DIR / cv_name)

        database.guardar_cv_usuario(user_id, cv_name)
        self.assertEqual(database.obtener_resumen_eventos_agente()["Pendiente"], 1)
        self.assertTrue(agent_worker.procesar_siguiente_evento("test-worker"))

        self.assertTrue(database.obtener_recomendaciones_usuario(user_id))
        event = database.obtener_eventos_agente(1)[0]
        self.assertEqual(event["estado"], "Completado")

    def test_worker_marks_unknown_event_as_error_after_last_retry(self):
        event_id = database.encolar_evento_agente("EVENTO_DESCONOCIDO", {}, max_intentos=1)
        self.assertTrue(agent_worker.procesar_siguiente_evento("test-worker"))
        event = database._rows("SELECT * FROM eventos_agente WHERE id=?", (event_id,))[0]
        self.assertEqual(event["estado"], "Error")
        self.assertIn("no soportado", event["ultimo_error"])

    def test_worker_detects_cv_copied_directly_into_uploads(self):
        document = Document()
        document.add_paragraph("Direct Upload")
        document.add_paragraph("direct@example.com")
        document.add_paragraph("Experiencia de 2 años en Python y SQL.")
        document.save(agent.UPLOADS_DIR / "direct_upload.docx")

        detected = agent_worker.detectar_cvs_nuevos()

        self.assertEqual(len(detected), 1)
        self.assertIn("direct_upload.docx", database.obtener_archivos_cv_registrados())
        self.assertEqual(database.obtener_resumen_eventos_agente()["Pendiente"], 1)

    def test_ollama_enriches_explanation_without_changing_deterministic_score(self):
        result = {
            "compatibilidad": 90.0, "estado": "Preseleccionado",
            "habilidades_encontradas": ["Python"], "habilidades_faltantes": ["SQL"],
            "desglose": {},
            "justificacion": "Base",
            "accion_sugerida": "Base",
        }
        response = {"message": {"content": """{
            "habilidades_adicionales": ["Docker"],
            "justificacion": "El perfil evidencia Python y Docker.",
            "accion_sugerida": "Validar SQL durante la entrevista.",
            "resumen_profesional": "Backend con experiencia en Python.",
            "fortalezas": ["Python", "Docker"],
            "brechas": ["SQL"],
            "preguntas_entrevista": ["¿Cómo desplegarías una API con Docker?"]
        }"""}}
        with patch.dict("os.environ", {"OLLAMA_ENABLED": "1"}, clear=False):
            with patch("modules.ollama_client._post_json", return_value=response):
                enriched = ollama_client.enriquecer_resultado("Python y Docker", {"titulo": "Backend"}, result)
        self.assertEqual(enriched["compatibilidad"], 90.0)
        self.assertIn("Docker", enriched["habilidades_encontradas"])
        self.assertIn("Ollama", enriched["desglose"]["motor"])
        self.assertEqual(enriched["brechas"], ["SQL"])

    def test_admin_can_open_candidate_intelligent_profile(self):
        vacancy = database.obtener_vacantes()[0]
        candidate_id = database.guardar_candidato({"nombre": "Ficha Demo", "correo": "ficha@example.com"})
        analysis_id = database.guardar_analisis(candidate_id, vacancy["id"], {
            "compatibilidad": 88, "estado": "Preseleccionado",
            "habilidades_encontradas": ["Python"], "habilidades_faltantes": ["SQL"],
            "resumen_profesional": "Perfil backend.",
            "fortalezas": ["Python"], "brechas": ["SQL"],
            "preguntas_entrevista": ["¿Cómo optimizarías una consulta SQL?"],
        })
        admin = app.test_client()
        self.login(admin, "admin@demo.com", "admin123")
        response = admin.get(f"/admin/ficha-candidato/{analysis_id}")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Perfil backend.", response.get_data(as_text=True))

    def test_high_score_creates_agent_alert(self):
        vacancy = database.obtener_vacantes()[0]
        candidate_id = database.guardar_candidato({"nombre": "Top Candidate", "correo": "top@example.com"})
        database.guardar_analisis(candidate_id, vacancy["id"], {
            "compatibilidad": 95, "estado": "Preseleccionado",
            "habilidades_encontradas": ["Python"], "habilidades_faltantes": [],
        })
        alerts = database.obtener_alertas_agente()
        self.assertEqual(alerts[0]["tipo"], "Candidato destacado")
        self.assertIn("95", alerts[0]["detalle"])

    def test_admin_can_retry_failed_worker_event(self):
        event_id = database.encolar_evento_agente("EVENTO_DESCONOCIDO", {}, max_intentos=1)
        agent_worker.procesar_siguiente_evento("test-worker")
        self.assertEqual(database.obtener_resumen_eventos_agente()["Error"], 1)
        self.assertTrue(database.obtener_alertas_agente())

        admin = app.test_client()
        self.login(admin, "admin@demo.com", "admin123")
        token = self.csrf(admin, "/admin/agente")
        response = admin.post(
            f"/admin/agente/eventos/{event_id}/reintentar",
            data={"_csrf_token": token},
        )
        self.assertEqual(response.status_code, 302)
        event = database._rows("SELECT * FROM eventos_agente WHERE id=?", (event_id,))[0]
        self.assertEqual(event["estado"], "Pendiente")
        self.assertEqual(event["intentos"], 0)

    def test_rrhh_decision_is_separate_from_ai_recommendation(self):
        vacancy = database.obtener_vacantes()[0]
        user_id = database.crear_usuario("Candidate", "candidate@example.com", "clave123")
        candidate_id = database.guardar_candidato({
            "nombre": "Candidate", "correo": "candidate@example.com", "usuario_id": user_id,
        })
        database.guardar_postulacion(user_id, vacancy["id"])
        database.guardar_analisis(candidate_id, vacancy["id"], {
            "compatibilidad": 95, "estado": "Preseleccionado",
            "habilidades_encontradas": ["Python"], "habilidades_faltantes": [],
        })
        database.actualizar_postulacion(user_id, vacancy["id"], 95, "Preseleccionado")

        agent.preparar_acciones_vacante(vacancy["id"])
        self.assertEqual(database.obtener_borradores_correo(vacancy["id"]), [])
        postulation = database.obtener_postulaciones_usuario(user_id)[0]
        self.assertEqual(postulation["estado"], "En revision por RRHH")
        self.assertEqual(postulation["recomendacion_ia"], "Preseleccionado")

        database.guardar_decision_rrhh(candidate_id, vacancy["id"], "Aprobado", "Perfil validado")
        agent.preparar_acciones_vacante(vacancy["id"])
        self.assertEqual(database.obtener_borradores_correo(vacancy["id"]), [])
        email_events = database._rows("SELECT * FROM eventos_agente WHERE tipo='ENVIAR_CORREOS' AND estado='Pendiente'")
        self.assertEqual(len(email_events), 1)
        self.assertEqual(database.obtener_invitaciones_usuario("candidate@example.com"), [])
        self.assertEqual(database.obtener_postulaciones_usuario(user_id)[0]["estado"], "Entrevista aprobada por RRHH")

    def test_approved_candidate_email_is_sent_automatically_by_worker(self):
        vacancy = database.obtener_vacantes()[0]
        candidate_id = database.guardar_candidato({"nombre": "Auto Mail", "correo": "auto@example.com"})
        database.guardar_analisis(candidate_id, vacancy["id"], {
            "compatibilidad": 91, "estado": "Preseleccionado",
            "habilidades_encontradas": ["Python"], "habilidades_faltantes": [],
        })

        database.guardar_decision_rrhh(candidate_id, vacancy["id"], "Aprobado", "Enviar invitacion")
        agent.preparar_acciones_vacante(vacancy["id"])
        queued = database._rows("SELECT * FROM correos WHERE destinatario=?", ("auto@example.com",))
        self.assertEqual(queued[0]["estado"], "En cola")

        with patch("modules.email_sender.enviar_correo_smtp") as smtp:
            self.assertTrue(agent_worker.procesar_siguiente_evento("auto-smtp-worker"))
        smtp.assert_called_once()
        sent = database._rows("SELECT estado FROM correos WHERE id=?", (queued[0]["id"],))[0]
        self.assertEqual(sent["estado"], "Enviado")

    def test_approved_emails_are_sent_by_worker_queue(self):
        vacancy = database.obtener_vacantes()[0]
        candidate_id = database.guardar_candidato({"nombre": "Demo", "correo": "demo@example.com"})
        email_id = database.guardar_correo(candidate_id, "demo@example.com", "Entrevista", "Mensaje", "Borrador", vacancy["id"])

        self.assertEqual(database.aprobar_borradores_correo(vacancy["id"]), 1)
        with patch("modules.email_sender.enviar_correo_smtp") as smtp:
            self.assertTrue(agent_worker.procesar_siguiente_evento("smtp-worker"))
        smtp.assert_called_once_with("demo@example.com", "Entrevista", "Mensaje")
        email_status = database._rows("SELECT estado FROM correos WHERE id=?", (email_id,))[0]["estado"]
        self.assertEqual(email_status, "Enviado")

    def test_worker_lease_can_be_renewed_only_by_owner(self):
        event_id = database.encolar_evento_agente("EVENTO_DESCONOCIDO", {})
        database.reclamar_evento_agente("worker-owner", lease_seconds=1)
        self.assertFalse(database.renovar_bloqueo_evento(event_id, "other-worker"))
        self.assertTrue(database.renovar_bloqueo_evento(event_id, "worker-owner", lease_seconds=120))

    def test_ollama_prompt_omits_contact_data(self):
        result = {
            "compatibilidad": 90.0, "estado": "Preseleccionado",
            "habilidades_encontradas": ["Python"], "habilidades_faltantes": [], "desglose": {},
        }
        response = {"message": {"content": """{
            "habilidades_adicionales": [], "justificacion": "Perfil válido.", "accion_sugerida": "Entrevistar.",
            "resumen_profesional": "Backend.", "fortalezas": ["Python"], "brechas": [],
            "preguntas_entrevista": ["Pregunta"]
        }"""}}
        with patch.dict("os.environ", {"OLLAMA_ENABLED": "1"}, clear=False):
            with patch("modules.ollama_client._post_json", return_value=response) as post_json:
                ollama_client.enriquecer_resultado(
                    "Correo person@example.com telefono +51 999 888 777\nDireccion: Calle 1\nPython",
                    {"titulo": "Backend"}, result,
                )
        prompt = post_json.call_args.args[1]["messages"][0]["content"]
        self.assertNotIn("person@example.com", prompt)
        self.assertNotIn("999 888 777", prompt)
        self.assertNotIn("Calle 1", prompt)

    def test_worker_omits_duplicate_cv_hash(self):
        document = Document()
        document.add_paragraph("Duplicated Candidate")
        document.add_paragraph("Experiencia de 2 años en Python y SQL.")
        first = agent.UPLOADS_DIR / "first.docx"
        second = agent.UPLOADS_DIR / "second.docx"
        document.save(first)
        shutil.copyfile(first, second)
        database.guardar_candidato({
            "nombre": "First", "archivo": first.name, "archivo_hash": calcular_hash_archivo(first),
        })
        self.assertEqual(agent_worker.detectar_cvs_nuevos(), [])
        self.assertFalse(any(item["archivo"] == second.name for item in database.obtener_candidatos()))


if __name__ == "__main__":
    unittest.main()
