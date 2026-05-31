"""Perception module: reads CV files and extracts a structured profile."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import fitz
import spacy
from docx import Document

from utils import normalize_text, unique

SKILL_CATALOG = [
    "Python", "Java", "JavaScript", "TypeScript", "C++", "C#", "PHP",
    "SQL", "PostgreSQL", "MySQL", "SQLite", "MongoDB", "Redis",
    "HTML", "CSS", "React", "Angular", "Vue", "Node.js", "Django",
    "Flask", "FastAPI", "Streamlit", "Spring", ".NET", "Laravel",
    "Git", "GitHub", "Docker", "Kubernetes", "Linux", "AWS", "Azure",
    "GCP", "Power BI", "Tableau", "Excel", "Pandas", "NumPy",
    "Scikit-Learn", "TensorFlow", "PyTorch", "spaCy", "NLP",
    "Machine Learning", "Deep Learning", "REST", "Scrum", "Agile",
    "Figma", "Selenium", "OpenCV", "R", "MATLAB", "SAP",
]

LANGUAGE_CATALOG = [
    "español", "ingles", "inglés", "english", "portugues", "portugués",
    "frances", "francés", "quechua", "aymara", "italiano", "aleman", "alemán",
]

CAREER_PATTERNS = [
    r"(?:ingenier[íi]a|licenciatura|bachiller|t[ée]cnico)[^\n,.]{0,80}",
    r"(?:administraci[óo]n|contabilidad|econom[íi]a|marketing|derecho)[^\n,.]{0,60}",
]


class CVExtractor:
    """Extracts text and likely CV fields with deterministic NLP rules."""

    def __init__(self) -> None:
        try:
            self.nlp = spacy.load("es_core_news_sm")
            self.model_name = "es_core_news_sm"
        except OSError:
            self.nlp = spacy.blank("es")
            self.model_name = "spacy.blank('es')"

    def extract(self, file_path: str | Path) -> dict[str, Any]:
        file_path = Path(file_path)
        text = self._read_file(file_path)
        if not text.strip():
            raise ValueError(f"No se pudo extraer texto de {file_path.name}")
        return {
            "file_name": file_path.name,
            "name": self._extract_name(text, file_path.stem),
            "email": self._first_match(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text),
            "phone": self._first_match(r"(?:\+?51[\s-]?)?(?:9\d{8}|\d{3}[\s-]\d{3}[\s-]\d{3})", text),
            "career": self._extract_career(text),
            "experience_years": self._extract_experience(text),
            "skills": self._catalog_matches(text, SKILL_CATALOG),
            "languages": self._catalog_matches(text, LANGUAGE_CATALOG),
            "certifications": self._section_items(text, ("certificaciones", "certificados", "cursos")),
            "projects": self._section_items(text, ("proyectos", "proyectos destacados")),
            "raw_text": text,
        }

    @staticmethod
    def _read_file(file_path: Path) -> str:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            with fitz.open(file_path) as document:
                return "\n".join(page.get_text("text") for page in document)
        if suffix == ".docx":
            document = Document(file_path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            return "\n".join(paragraphs + table_cells)
        raise ValueError("Formato no soportado. Use archivos PDF o DOCX.")

    def _extract_name(self, text: str, fallback: str) -> str:
        doc = self.nlp(text[:1500])
        people = [entity.text.strip() for entity in doc.ents if entity.label_ == "PER"]
        if people:
            return people[0]
        for line in text.splitlines()[:8]:
            cleaned = line.strip(" -:\t")
            if (
                2 <= len(cleaned.split()) <= 5
                and len(cleaned) <= 70
                and "@" not in cleaned
                and not any(char.isdigit() for char in cleaned)
                and normalize_text(cleaned) not in {"curriculum vitae", "hoja de vida", "cv"}
            ):
                return cleaned.title()
        return fallback.replace("_", " ").title()

    @staticmethod
    def _extract_career(text: str) -> str:
        for pattern in CAREER_PATTERNS:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                return match.group(0).strip().title()
        return "No identificada"

    @staticmethod
    def _extract_experience(text: str) -> float:
        patterns = [
            r"(\d+(?:[.,]\d+)?)\s*(?:años|anos)\s+(?:de\s+)?experiencia",
            r"experiencia[^\n]{0,30}?(\d+(?:[.,]\d+)?)\s*(?:años|anos)",
        ]
        values: list[float] = []
        for pattern in patterns:
            values.extend(
                float(value.replace(",", "."))
                for value in re.findall(pattern, text, re.IGNORECASE)
            )
        return max(values, default=0.0)

    @staticmethod
    def _catalog_matches(text: str, catalog: list[str]) -> list[str]:
        normalized = normalize_text(text)
        found = []
        for item in catalog:
            token = normalize_text(item)
            if re.search(rf"(?<!\w){re.escape(token)}(?!\w)", normalized):
                found.append(item)
        return unique(found)

    @staticmethod
    def _section_items(text: str, headings: tuple[str, ...]) -> list[str]:
        heading_group = "|".join(re.escape(item) for item in headings)
        match = re.search(
            rf"(?:{heading_group})\s*:?\s*\n(.+?)(?=\n[A-ZÁÉÍÓÚÑ ]{{4,}}\s*:?\s*\n|\Z)",
            text,
            re.IGNORECASE | re.DOTALL,
        )
        if not match:
            return []
        lines = [
            line.strip(" -*•\t")
            for line in match.group(1).splitlines()
            if line.strip(" -*•\t")
        ]
        return unique(lines[:8])

    @staticmethod
    def _first_match(pattern: str, text: str) -> str:
        match = re.search(pattern, text, re.IGNORECASE)
        return match.group(0).strip() if match else ""
