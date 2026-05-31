"""Unit tests for the complete autonomous workflow."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from sys import path

path.insert(0, str(Path(__file__).resolve().parents[1]))

from analyzer import CandidateAnalyzer
from database import Database
from email_sender import SMTPConfig, build_invitation, send_email
from extractor import CVExtractor
from ranking import rank_candidates, select_top_candidates
from reports import generate_excel, generate_pdf


class AgentWorkflowTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.temp_dir.name) / "test.db")
        self.vacancy = {
            "title": "Desarrollador Python",
            "area": "Tecnología",
            "description": "Desarrollo de aplicaciones y análisis inteligente.",
            "requirements": "Conocimientos de desarrollo, datos y control de versiones.",
            "required_experience": 2,
            "required_skills": ["Python", "SQL", "Git", "Streamlit"],
        }

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_autonomous_workflow(self) -> None:
        vacancy_id = self.db.create_vacancy(self.vacancy)
        vacancy = self.db.get_vacancy(vacancy_id)
        process_id = self.db.create_process(vacancy_id)
        analyzer = CandidateAnalyzer()
        profiles = [
            self._profile("Ana Torres", "ana@example.com", "Python SQL Git Streamlit", 3),
            self._profile("Bruno Flores", "bruno@example.com", "Python SQL Git", 2),
            self._profile("Carla Díaz", "carla@example.com", "Excel Power BI", 1),
            self._profile("Diego Ramos", "diego@example.com", "Python Git", 1),
        ]
        for profile in profiles:
            analysis = analyzer.analyze(vacancy, profile)
            self.db.save_candidate(process_id, {"profile": profile, "analysis": analysis})
        self.db.update_process_summary(process_id)
        candidates = select_top_candidates(self.db.list_candidates(process_id))
        self.db.update_selection(process_id, [item["id"] for item in candidates if item["selected"]])

        ranked = rank_candidates(self.db.list_candidates(process_id))
        self.assertEqual(ranked[0]["name"], "Ana Torres")
        self.assertEqual(sum(item["selected"] for item in ranked), 3)
        self.assertGreater(ranked[0]["score"], ranked[-1]["score"])
        self.assertGreater(len(generate_excel(vacancy, ranked)), 1000)
        self.assertGreater(len(generate_pdf(vacancy, ranked)), 1000)

    def test_docx_extraction_and_simulated_email(self) -> None:
        from docx import Document

        cv_path = Path(self.temp_dir.name) / "ana.docx"
        document = Document()
        document.add_paragraph("Ana Torres Mendoza")
        document.add_paragraph("ana@example.com | 987654321")
        document.add_paragraph("Ingeniería de Sistemas. 3 años de experiencia.")
        document.add_paragraph("Python, SQL, Git, Streamlit")
        document.save(cv_path)
        profile = CVExtractor().extract(cv_path)
        subject, body = build_invitation(profile["name"], "Desarrollador Python", "2026-06-10", "09:00")
        result = send_email(SMTPConfig("", "", simulation_mode=True), profile["email"], subject, body)

        self.assertEqual(profile["email"], "ana@example.com")
        self.assertEqual(profile["experience_years"], 3)
        self.assertIn("Python", profile["skills"])
        self.assertEqual(result["status"], "SIMULADO")

    @staticmethod
    def _profile(name: str, email: str, text: str, years: float) -> dict:
        return {
            "file_name": f"{name}.docx",
            "name": name,
            "email": email,
            "phone": "",
            "career": "Ingeniería de Sistemas",
            "experience_years": years,
            "skills": text.split(),
            "languages": ["Español"],
            "certifications": [],
            "projects": [],
            "raw_text": text,
        }


if __name__ == "__main__":
    unittest.main()
