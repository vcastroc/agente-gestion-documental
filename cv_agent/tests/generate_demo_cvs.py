"""Generate DOCX CVs for the classroom demonstration."""

from pathlib import Path

from docx import Document

BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = BASE_DIR / "data" / "demo_cvs"

CVS = [
    {
        "filename": "ana_torres.docx",
        "name": "Ana Torres Mendoza",
        "email": "ana.torres@example.com",
        "phone": "987654321",
        "career": "Ingeniería de Sistemas",
        "experience": "3 años de experiencia",
        "skills": "Python, SQL, Git, Streamlit, Scikit-Learn, Pandas, Docker",
        "languages": "Español, Inglés",
        "projects": "Dashboard de selección de personal con Streamlit y SQLite",
    },
    {
        "filename": "bruno_quinto.docx",
        "name": "Bruno Quinto Flores",
        "email": "bruno.quinto@example.com",
        "phone": "976543210",
        "career": "Ingeniería Informática",
        "experience": "2 años de experiencia",
        "skills": "Python, SQL, Git, Flask, Pandas",
        "languages": "Español, Inglés",
        "projects": "API REST para inventario universitario",
    },
    {
        "filename": "carla_rojas.docx",
        "name": "Carla Rojas Díaz",
        "email": "carla.rojas@example.com",
        "phone": "965432109",
        "career": "Administración",
        "experience": "1 año de experiencia",
        "skills": "Excel, Power BI, Scrum",
        "languages": "Español",
        "projects": "Reporte comercial con Power BI",
    },
]


def create_cv(data: dict[str, str]) -> None:
    document = Document()
    document.add_heading(data["name"], 0)
    document.add_paragraph(f"Correo: {data['email']}")
    document.add_paragraph(f"Teléfono: {data['phone']}")
    document.add_heading("Perfil", level=1)
    document.add_paragraph(f"{data['career']}. {data['experience']}.")
    document.add_heading("Habilidades", level=1)
    document.add_paragraph(data["skills"])
    document.add_heading("Idiomas", level=1)
    document.add_paragraph(data["languages"])
    document.add_heading("Proyectos", level=1)
    document.add_paragraph(data["projects"])
    document.save(OUTPUT_DIR / data["filename"])


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for cv in CVS:
        create_cv(cv)
    print(f"CVs de demostración generados en: {OUTPUT_DIR}")
